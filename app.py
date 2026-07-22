from flask import Flask, render_template, request, send_file, Response, stream_with_context
import json
import os
import io
from PIL import Image
import pytesseract
from pdf2image import convert_from_path, pdfinfo_from_path
from docx import Document

import platform

if platform.system() == "Windows":
    pytesseract.pytesseract.tesseract_cmd = r"C:\Program Files\Tesseract-OCR\tesseract.exe"
    POPPLER_PATH = (
        r"C:\Users\MANU M M\AppData\Local\Microsoft\WinGet\Packages"
        r"\oschwartz10612.Poppler_Microsoft.Winget.Source_8wekyb3d8bbwe"
        r"\poppler-25.07.0\Library\bin"
    )
else:
    # Linux (Docker / Render)
    pytesseract.pytesseract.tesseract_cmd = "/usr/bin/tesseract"
    POPPLER_PATH = "/usr/bin"

app = Flask(__name__, static_folder="static", template_folder="templates")
UPLOAD_FOLDER = "uploads"
MAX_PAGES_PER_REQUEST = 20
os.makedirs(UPLOAD_FOLDER, exist_ok=True)


def parse_page_selection(page_selection):
    page_selection = page_selection.strip()

    if not page_selection or page_selection.lower() == "all":
        return None

    pages = set()

    try:
        for part in page_selection.split(","):
            part = part.strip()

            if "-" in part:
                start, end = part.split("-", 1)
                start = int(start)
                end = int(end)

                if start > end:
                    raise ValueError

                pages.update(range(start, end + 1))
            else:
                pages.add(int(part))

        return sorted(pages)

    except ValueError:
        raise ValueError(
            "Invalid page selection. Examples: All, 3, 1-3, 2,5,8"
        )


@app.route("/")
def home():
    return render_template(
        "index.html",
        status="Ready",
        extracted_text="",
        selected_language="auto",
        other_language=""
    )


@app.route("/upload", methods=["POST"])
def upload():
    if "file" not in request.files or request.files["file"].filename == "":
        return render_template(
            "index.html",
            status="No file selected.",
            extracted_text=""
        )

    lang = request.form.get("language", "auto")
    other = request.form.get("other_language", "").strip()
    try:
        page_selection = parse_page_selection(
            request.form.get("pageSelection", "All")
        )
    except ValueError as e:
        return render_template(
            "index.html",
            status=str(e),
            extracted_text="",
            selected_language=lang,
            other_language=other
        )

    if lang == "other":
        ocr_lang = other or "eng"
    elif lang == "auto":
        ocr_lang = None
    else:
        ocr_lang = lang

    f = request.files["file"]
    path = os.path.join(UPLOAD_FOLDER, f.filename)
    f.save(path)

    ext = os.path.splitext(path)[1].lower()

    try:
        if ext == ".pdf":
            pages = convert_from_path(path, poppler_path=POPPLER_PATH)
            out = []
            if page_selection is not None:
                total_pages = len(pages)

                if max(page_selection) > total_pages or min(page_selection) < 1:
                    raise ValueError(
                        f"Invalid page selection. This PDF has {total_pages} pages."
                    )

            for i, p in enumerate(pages, 1):
                if page_selection is not None and i not in page_selection:
                    continue
                out.append(f"\nPage {i}\n\n")

                if ocr_lang:
                    out.append(
                        pytesseract.image_to_string(
                            p,
                            lang=ocr_lang
                        )
                    )
                else:
                    out.append(
                        pytesseract.image_to_string(p)
                    )

                out.append("\n")

            text = "".join(out)

        else:
            img = Image.open(path)

            if ocr_lang:
                text = pytesseract.image_to_string(
                    img,
                    lang=ocr_lang
                )
            else:
                text = pytesseract.image_to_string(img)

        if ext == ".pdf":
            status = f"OCR completed successfully. {len(pages)} page(s) processed."
        else:
            status = "OCR completed successfully. Image processed."
    except Exception as e:
        import traceback
        traceback.print_exc()

        status = f"OCR failed: {e}"
        text = f"{type(e).__name__}: {e}"

    return render_template(
        "index.html",
        status=status,
        extracted_text=text,
        selected_language=lang,
        other_language=other
    )


@app.route("/download-docx", methods=["POST"])
def download_docx():
    text = request.form.get("text", "")

    document = Document()
    document.add_heading("OCR Result", level=1)
    document.add_paragraph(text)

    memory_file = io.BytesIO()
    document.save(memory_file)
    memory_file.seek(0)

    return send_file(
        memory_file,
        as_attachment=True,
        download_name="ocr_result.docx",
        mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )


def generate_ocr_stream(path, ext, page_selection, ocr_lang):
    if ext == ".pdf":
        info = pdfinfo_from_path(path, poppler_path=POPPLER_PATH)
        total = info["Pages"]

        if page_selection is not None:
            if max(page_selection) > total or min(page_selection) < 1:
                yield json.dumps({
                    "type": "error",
                    "message": f"Invalid page selection. This PDF has {total} pages."
                }) + "\n"
                return
            indices_all = sorted(page_selection)
        else:
            indices_all = list(range(1, total + 1))

        if len(indices_all) > MAX_PAGES_PER_REQUEST:
            batch = indices_all[:MAX_PAGES_PER_REQUEST]
            remaining = indices_all[MAX_PAGES_PER_REQUEST:]
        else:
            batch = indices_all
            remaining = []

        yield json.dumps({"type": "start", "total": len(batch)}) + "\n"

        for count, i in enumerate(batch, 1):
            try:
                page_images = convert_from_path(
                    path, poppler_path=POPPLER_PATH,
                    first_page=i, last_page=i
                )
                p = page_images[0]

                if ocr_lang:
                    text = pytesseract.image_to_string(p, lang=ocr_lang)
                else:
                    text = pytesseract.image_to_string(p)
            except Exception as e:
                yield json.dumps({
                    "type": "error",
                    "message": f"Page {i} failed: {e}"
                }) + "\n"
                continue

            yield json.dumps({
                "type": "page",
                "page": i,
                "index": count,
                "total": len(batch),
                "text": text
            }) + "\n"

        if remaining:
            next_range = f"{remaining[0]}-{remaining[-1]}"
            yield json.dumps({
                "type": "done",
                "message": (
                    f"Processed pages {batch[0]}-{batch[-1]} of {total}. "
                    f"{len(remaining)} page(s) remaining."
                ),
                "more": True,
                "next_range": next_range
            }) + "\n"
        else:
            yield json.dumps({
                "type": "done",
                "message": f"OCR completed successfully. {len(batch)} page(s) processed.",
                "more": False
            }) + "\n"

    else:
        try:
            img = Image.open(path)
            text = (
                pytesseract.image_to_string(img, lang=ocr_lang)
                if ocr_lang else pytesseract.image_to_string(img)
            )
            yield json.dumps({
                "type": "page", "page": 1, "index": 1, "total": 1, "text": text
            }) + "\n"
            yield json.dumps({
                "type": "done", "message": "OCR completed successfully. Image processed."
            }) + "\n"
        except Exception as e:
            yield json.dumps({"type": "error", "message": str(e)}) + "\n"


@app.route("/upload-stream", methods=["POST"])
def upload_stream():
    if "file" not in request.files or request.files["file"].filename == "":
        return Response(
            json.dumps({"type": "error", "message": "No file selected."}) + "\n",
            mimetype="application/x-ndjson"
        )

    lang = request.form.get("language", "auto")
    other = request.form.get("other_language", "").strip()

    try:
        page_selection = parse_page_selection(
            request.form.get("pageSelection", "All")
        )
    except ValueError as e:
        return Response(
            json.dumps({"type": "error", "message": str(e)}) + "\n",
            mimetype="application/x-ndjson"
        )

    if lang == "other":
        ocr_lang = other or "eng"
    elif lang == "auto":
        ocr_lang = None
    else:
        ocr_lang = lang

    f = request.files["file"]
    path = os.path.join(UPLOAD_FOLDER, f.filename)
    f.save(path)
    ext = os.path.splitext(path)[1].lower()

    return Response(
        stream_with_context(
            generate_ocr_stream(path, ext, page_selection, ocr_lang)
        ),
        mimetype="application/x-ndjson"
    )


if __name__ == "__main__":
    app.run(host="0.0.0.0", port=7860, debug=True)